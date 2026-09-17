from io import BytesIO

from django.contrib.auth.models import Group, User
from openpyxl import load_workbook
from rest_framework.test import APIClient, APITestCase

from app.models import (
    CategoriaDeSalgado,
    Conta,
    Desperdicio,
    Encarregado,
    FechamentoCaixa,
    Salgado,
)
from app.services.catalogo_de_salgados import garantir_catalogo_de_salgados

from .fabricas import DIA_COMUM, criar_loja


class BaseDoDesperdicio(APITestCase):
    """So o cenario: catalogo semeado, loja, aparelho autenticado.

    Sem teste nenhum aqui de proposito — herdar de uma classe QUE TEM testes
    faria o runner repetir todos eles dentro de cada filha.
    """

    def setUp(self):
        from django.core.cache import cache

        cache.clear()
        self.client = APIClient()
        self.conta = Conta.objects.create(nome="Marina")
        self.loja = criar_loja(
            nome_loja="Centro", cidade="Patos", endereco="Rua 1", conta=self.conta
        )
        garantir_catalogo_de_salgados(self.conta)
        self.coxinha = Salgado.objects.get(
            categoria__conta=self.conta,
            categoria__nome="Salgados grande",
            nome="Coxinha",
        )
        self.kibe = Salgado.objects.get(
            categoria__conta=self.conta,
            categoria__nome="Salgados grande",
            nome="Kibe",
        )
        # O brief nao cria encarregado nenhum, mas `lancado_por` e obrigatorio
        # no serializer de escrita (mesmo padrao de test_consumo.py) — sem
        # ele todo POST de fechamento cai com 400 antes de chegar perto do
        # desperdicio.
        self.ana = Encarregado.objects.create(nome="Ana Paula", conta=self.conta)
        self.entrar()

    def entrar(self, conta=None):
        conta = conta or self.conta
        self.client.post(
            "/api/v1/formulario/acesso/",
            {"codigo": conta.codigo_acesso, "apelido": "Caixa"},
            format="json",
        )

    def payload(self, **extra):
        base = {
            "loja": str(self.loja.public_id),
            "lancado_por": str(self.ana.public_id),
            "data": DIA_COMUM,
            "periodo": "TARDE",
            "pix": "100.00",
            "cartao": "0.00",
            "dinheiro": "50.00",
            "link_pagamento": "0.00",
            "houve_retirada": False,
            "houve_desperdicio": False,
        }
        base.update(extra)
        return base

class DesperdicioNoEnvioTests(BaseDoDesperdicio):
    def test_lanca_duas_linhas_de_desperdicio(self):
        resp = self.client.post(
            "/api/v1/fechamentos-caixa/",
            self.payload(
                desperdicios=[
                    {"salgado": str(self.coxinha.public_id), "quantidade": 8},
                    {"salgado": str(self.kibe.public_id), "quantidade": 3},
                ]
            ),
            format="json",
        )

        self.assertEqual(resp.status_code, 201)
        self.assertEqual(Desperdicio.objects.count(), 2)

    def test_o_desperdicio_nao_mexe_no_total(self):
        """Ninguem pagou nada: nenhum dinheiro deixou a gaveta. Somar aqui
        faria o fechamento acusar uma diferenca que nao existe."""
        resp = self.client.post(
            "/api/v1/fechamentos-caixa/",
            self.payload(
                desperdicios=[
                    {"salgado": str(self.coxinha.public_id), "quantidade": 40}
                ]
            ),
            format="json",
        )

        self.assertEqual(resp.status_code, 201)
        self.assertEqual(str(resp.data["total"]), "150.00")

    def test_quantidade_zero_e_recusada(self):
        resp = self.client.post(
            "/api/v1/fechamentos-caixa/",
            self.payload(
                desperdicios=[
                    {"salgado": str(self.coxinha.public_id), "quantidade": 0}
                ]
            ),
            format="json",
        )

        self.assertEqual(resp.status_code, 400)
        self.assertEqual(FechamentoCaixa.objects.count(), 0)

    def test_salgado_de_outra_empresa_e_recusado(self):
        """O endpoint e publico: nada impede colar no payload o id de um item
        da empresa vizinha."""
        outra = Conta.objects.create(nome="Primavera")
        garantir_catalogo_de_salgados(outra)
        alheio = Salgado.objects.filter(categoria__conta=outra).first()

        resp = self.client.post(
            "/api/v1/fechamentos-caixa/",
            self.payload(
                desperdicios=[{"salgado": str(alheio.public_id), "quantidade": 2}]
            ),
            format="json",
        )

        self.assertEqual(resp.status_code, 400)

    def test_salgado_inativo_fica_fora(self):
        self.coxinha.ativo = False
        self.coxinha.save()

        resp = self.client.post(
            "/api/v1/fechamentos-caixa/",
            self.payload(
                desperdicios=[
                    {"salgado": str(self.coxinha.public_id), "quantidade": 2}
                ]
            ),
            format="json",
        )

        self.assertEqual(resp.status_code, 400)

    def test_categoria_inativa_tira_o_item_da_lista(self):
        categoria = self.coxinha.categoria
        categoria.ativo = False
        categoria.save()

        resp = self.client.post(
            "/api/v1/fechamentos-caixa/",
            self.payload(
                desperdicios=[
                    {"salgado": str(self.coxinha.public_id), "quantidade": 2}
                ]
            ),
            format="json",
        )

        self.assertEqual(resp.status_code, 400)

    def test_duas_linhas_do_mesmo_item_sao_aceitas(self):
        """Mesma escolha do Consumo: recusar obrigaria a loja a somar de cabeca
        antes de digitar."""
        resp = self.client.post(
            "/api/v1/fechamentos-caixa/",
            self.payload(
                desperdicios=[
                    {"salgado": str(self.coxinha.public_id), "quantidade": 3},
                    {"salgado": str(self.coxinha.public_id), "quantidade": 5},
                ]
            ),
            format="json",
        )

        self.assertEqual(resp.status_code, 201)
        self.assertEqual(Desperdicio.objects.count(), 2)

    def test_sem_desperdicio_o_fechamento_continua_igual(self):
        resp = self.client.post(
            "/api/v1/fechamentos-caixa/", self.payload(), format="json"
        )

        self.assertEqual(resp.status_code, 201)
        self.assertEqual(Desperdicio.objects.count(), 0)


class CorrecaoDoDesperdicioTests(BaseDoDesperdicio):
    """A correcao da loja, dentro dos 20 minutos.

    O que esta em jogo e a diferenca entre "o payload nao falou de
    desperdicio" (None) e "nao teve nenhum" (lista vazia). Sem ela, corrigir o
    PIX apagaria as perdas do turno de tabela — e ninguem perceberia, porque a
    correcao devolve 200 do mesmo jeito.
    """

    def enviar_com_perda(self):
        resp = self.client.post(
            "/api/v1/fechamentos-caixa/",
            self.payload(
                desperdicios=[
                    {"salgado": str(self.coxinha.public_id), "quantidade": 8}
                ]
            ),
            format="json",
        )
        self.assertEqual(resp.status_code, 201)
        return resp.data["id"]

    def test_corrigir_o_pix_nao_apaga_o_desperdicio(self):
        """A correcao exige o formulario inteiro (mesmo padrao do consumo em
        test_consumo.py): sem "desperdicios" no corpo e o payload nao falando
        do assunto, nao apagar o que ja estava la."""
        public_id = self.enviar_com_perda()
        payload = self.payload(pix="120.00")
        payload.pop("desperdicios", None)

        resp = self.client.patch(
            f"/api/v1/fechamentos-caixa/{public_id}/correcao/",
            payload,
            format="json",
        )

        self.assertEqual(resp.status_code, 200, resp.data)
        self.assertEqual(Desperdicio.objects.count(), 1)

    def test_lista_vazia_apaga_as_linhas(self):
        """"Nao teve desperdicio nenhum" tem que poder ser dito."""
        public_id = self.enviar_com_perda()

        resp = self.client.patch(
            f"/api/v1/fechamentos-caixa/{public_id}/correcao/",
            self.payload(desperdicios=[]),
            format="json",
        )

        self.assertEqual(resp.status_code, 200, resp.data)
        self.assertEqual(Desperdicio.objects.count(), 0)

    def test_lista_nova_substitui_a_antiga_sem_somar(self):
        public_id = self.enviar_com_perda()

        resp = self.client.patch(
            f"/api/v1/fechamentos-caixa/{public_id}/correcao/",
            self.payload(
                desperdicios=[
                    {"salgado": str(self.kibe.public_id), "quantidade": 2}
                ]
            ),
            format="json",
        )

        self.assertEqual(resp.status_code, 200, resp.data)
        self.assertEqual(Desperdicio.objects.count(), 1)
        self.assertEqual(Desperdicio.objects.first().salgado, self.kibe)

    def test_o_get_da_correcao_devolve_as_linhas_de_desperdicio(self):
        """O formulario se preenche de novo a partir daqui (ver docstring do
        FechamentoCaixaFormularioSerializer). Sem "desperdicios" nesta
        resposta, o front nao tem como reidratar o bloco — e a correcao de
        outro campo manda a lista vazia de volta, que apaga o que estava la."""
        public_id = self.enviar_com_perda()

        resp = self.client.get(f"/api/v1/fechamentos-caixa/{public_id}/correcao/")

        self.assertEqual(resp.status_code, 200, resp.data)
        self.assertEqual(len(resp.data["desperdicios"]), 1)
        linha = resp.data["desperdicios"][0]
        self.assertEqual(str(linha["salgado"]), str(self.coxinha.public_id))
        self.assertEqual(linha["nome"], "Coxinha")
        self.assertEqual(linha["quantidade"], 8)


class DesperdicioNaLeituraTests(APITestCase):
    def setUp(self):
        self.conta = Conta.objects.create(nome="Marina")
        self.loja = criar_loja(
            nome_loja="Centro", cidade="Patos", endereco="Rua 1", conta=self.conta
        )
        garantir_catalogo_de_salgados(self.conta)
        coxinha = Salgado.objects.get(
            categoria__conta=self.conta,
            categoria__nome="Salgados grande",
            nome="Coxinha",
        )
        fechamento = FechamentoCaixa.objects.create(
            loja=self.loja, nome_funcionario="Ana", data=DIA_COMUM, periodo="TARDE"
        )
        Desperdicio.objects.create(
            fechamento=fechamento, salgado=coxinha, quantidade=8
        )

        self.gerente = User.objects.create_user(
            username="gerente@marina.com", password="x"
        )
        Group.objects.get_or_create(name="Gerente")[0].user_set.add(self.gerente)
        from .fabricas import vincular_conta

        vincular_conta(self.gerente, self.conta)
        self.client = APIClient()
        self.client.force_authenticate(self.gerente)

    def test_o_painel_recebe_as_linhas_aninhadas(self):
        resp = self.client.get("/api/v1/fechamentos-caixa/")

        self.assertEqual(resp.status_code, 200)
        dados = resp.data["results"] if "results" in resp.data else resp.data
        linhas = dados[0]["desperdicios"]
        self.assertEqual(len(linhas), 1)
        self.assertEqual(linhas[0]["nome"], "Coxinha")
        self.assertEqual(linhas[0]["categoria_nome"], "Salgados grande")
        self.assertEqual(linhas[0]["quantidade"], 8)


class PlanilhaComDesperdicioTests(APITestCase):
    def setUp(self):
        self.conta = Conta.objects.create(nome="Marina")
        self.loja = criar_loja(
            nome_loja="Centro", cidade="Patos", endereco="Rua 1", conta=self.conta
        )
        garantir_catalogo_de_salgados(self.conta)
        coxinha = Salgado.objects.get(
            categoria__conta=self.conta,
            categoria__nome="Salgados grande",
            nome="Coxinha",
        )
        kibe = Salgado.objects.get(
            categoria__conta=self.conta,
            categoria__nome="Salgados grande",
            nome="Kibe",
        )
        fechamento = FechamentoCaixa.objects.create(
            loja=self.loja, nome_funcionario="Ana", data=DIA_COMUM, periodo="TARDE"
        )
        Desperdicio.objects.create(
            fechamento=fechamento, salgado=coxinha, quantidade=8
        )
        Desperdicio.objects.create(fechamento=fechamento, salgado=kibe, quantidade=3)

        self.gerente = User.objects.create_user(
            username="gerente@marina.com", password="x"
        )
        Group.objects.get_or_create(name="Gerente")[0].user_set.add(self.gerente)
        from .fabricas import vincular_conta

        vincular_conta(self.gerente, self.conta)
        self.client = APIClient()
        self.client.force_authenticate(self.gerente)

    def baixar(self):
        resp = self.client.get(f"/api/v1/planilha/?de={DIA_COMUM}&ate={DIA_COMUM}")
        self.assertEqual(resp.status_code, 200)
        return load_workbook(BytesIO(resp.content))

    def test_a_aba_existe(self):
        planilha = self.baixar()

        self.assertIn("Desperdício", planilha.sheetnames)

    def test_as_linhas_e_o_subtotal(self):
        folha = self.baixar()["Desperdício"]
        linhas = list(folha.values)

        # cabecalho + 2 linhas + subtotal da loja + TOTAL
        self.assertEqual(len(linhas), 5)
        quantidades = [linha[5] for linha in linhas[1:3]]
        self.assertEqual(sorted(quantidades), [3, 8])
        self.assertEqual(linhas[3][5], 11)  # Total Centro
        self.assertEqual(linhas[4][5], 11)  # TOTAL

    def test_a_quantidade_nao_e_formatada_como_dinheiro(self):
        """Quantidade na coluna com formato de dinheiro convidaria a somar
        unidades com reais — o erro que a separacao das abas existe para
        impedir."""
        folha = self.baixar()["Desperdício"]

        self.assertNotEqual(folha.cell(row=2, column=6).number_format, "#,##0.00")

    def test_a_aba_de_saidas_nao_ganhou_o_desperdicio(self):
        """Somar a aba Saidas tem que continuar dando o mesmo que as colunas
        de retirada, despesa e devolucao da aba Entradas."""
        folha = self.baixar()["Saídas"]

        tipos = {linha[3] for linha in list(folha.values)[1:] if linha[3]}
        self.assertNotIn("Desperdício", tipos)
