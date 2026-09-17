from datetime import timedelta
from io import BytesIO

from django.contrib.auth.models import Group, User
from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from app.models import Encarregado, FechamentoCaixa

from .fabricas import DIA_COMUM, conta_padrao, criar_loja, lista, vincular_conta


def loja_padrao():
    """Uma loja so para todos os fechamentos do teste.

    `criar_loja()` de fabricas.py cria uma loja NOVA a cada chamada, e dois
    fechamentos em lojas diferentes nao provariam nada sobre o contador da loja
    nem sobre o mesmo turno voltar a ficar livre.
    """
    from app.models import Loja

    loja, _ = Loja.objects.get_or_create(
        conta=conta_padrao(),
        nome_loja="Loja A",
        defaults={"cidade": "Patos", "endereco": "Rua 1"},
    )
    return loja


def criar_fechamento(**campos):
    campos.setdefault("loja", loja_padrao())
    campos.setdefault("data", DIA_COMUM)
    campos.setdefault("periodo", "MANHA")
    campos.setdefault("nome_funcionario", "Marina")
    # lancado_por e nulavel no modelo, entao fica de fora: quem lancou nao
    # muda nada do que este arquivo testa.
    return FechamentoCaixa.objects.create(**campos)


class ManagerDeCanceladosTests(TestCase):
    def setUp(self):
        conta_padrao()
        self.vivo = criar_fechamento()
        self.morto = criar_fechamento(periodo="TARDE")
        self.morto.cancelado_em = timezone.now()
        self.morto.save(update_fields=["cancelado_em"])

    def test_o_padrao_nao_enxerga_o_cancelado(self):
        """E daqui que sai o painel, a planilha, os graficos e as saidas."""
        ids = set(FechamentoCaixa.objects.values_list("id", flat=True))
        self.assertEqual(ids, {self.vivo.id})

    def test_todos_enxerga_os_dois(self):
        """O /admin e a propria acao de cancelar precisam ver o cancelado."""
        self.assertEqual(FechamentoCaixa.todos.count(), 2)

    def test_a_relacao_da_loja_continua_enxergando_o_cancelado(self):
        """Apagar loja com historico continua barrado: o lancamento cancelado
        ainda e historico, e some junto se a loja sair."""
        self.assertEqual(self.vivo.loja.fechamentos_caixa.count(), 2)

    def test_a_property_diz_o_estado(self):
        self.assertFalse(self.vivo.cancelado)
        self.assertTrue(FechamentoCaixa.todos.get(id=self.morto.id).cancelado)

    def test_nasce_nao_cancelado(self):
        self.assertIsNone(self.vivo.cancelado_em)
        self.assertIsNone(self.vivo.cancelado_por)


class AdminEnxergaCanceladoTests(TestCase):
    def test_a_lista_do_admin_mostra_o_cancelado(self):
        """Sem isto o dono da plataforma perde de vista o que foi cancelado —
        justamente o registro que o cancelamento existe para preservar."""
        conta_padrao()
        fechamento = criar_fechamento()
        fechamento.cancelado_em = timezone.now()
        fechamento.save(update_fields=["cancelado_em"])

        admin = User.objects.create_superuser("dono@x.com", "dono@x.com", "123456")
        self.client.force_login(admin)
        resposta = self.client.get("/admin/app/fechamentocaixa/")

        self.assertEqual(resposta.status_code, 200)
        self.assertContains(resposta, "Loja A")


class PlanilhaIgnoraCanceladoTests(TestCase):
    """Prova que a exportacao para Excel nao arrasta o cancelado junto.

    Gera o .xlsx de verdade pelo endpoint publico e le a aba "Entradas" de
    volta com openpyxl — o mesmo caminho que test_planilha.py ja usa — porque
    so checar o 200 da resposta nao provaria que a LINHA do cancelado sumiu.
    """

    def test_o_cancelado_nao_entra_na_planilha(self):
        for nome in ("Admin", "Gerente"):
            Group.objects.get_or_create(name=nome)
        conta = conta_padrao()
        vivo = criar_fechamento(nome_funcionario="Marina")
        morto = criar_fechamento(periodo="TARDE", nome_funcionario="Paulo")
        morto.cancelado_em = timezone.now()
        morto.save(update_fields=["cancelado_em"])

        gerente = User.objects.create_user(username="ger-planilha@x.com", password="123456")
        gerente.groups.add(Group.objects.get(name="Gerente"))
        vincular_conta(gerente, conta)

        client = APIClient()
        client.force_authenticate(gerente)
        resposta = client.get(f"/api/v1/planilha/?de={DIA_COMUM}&ate={DIA_COMUM}")
        self.assertEqual(resposta.status_code, 200, resposta.content[:200])

        planilha = load_workbook(BytesIO(resposta.content))
        entradas = planilha["Entradas"]
        nomes = [
            linha[3] for linha in entradas.iter_rows(min_row=2, values_only=True)
        ]
        self.assertIn("Marina", nomes)
        self.assertNotIn("Paulo", nomes)


class CancelarPelaGerenciaTests(TestCase):
    def setUp(self):
        for nome in ("Admin", "Gerente", "Funcionario"):
            Group.objects.get_or_create(name=nome)
        self.conta = conta_padrao()
        self.fechamento = criar_fechamento()

        self.gerente = User.objects.create_user(username="ger@x.com", password="123456")
        self.gerente.groups.add(Group.objects.get(name="Gerente"))
        vincular_conta(self.gerente, self.conta)

        self.client_api = APIClient()
        self.client_api.force_authenticate(self.gerente)

    def url(self):
        return f"/api/v1/fechamentos-caixa/{self.fechamento.public_id}/cancelar/"

    def test_gerencia_cancela_e_o_lancamento_some_da_listagem(self):
        resposta = self.client_api.post(self.url())
        self.assertEqual(resposta.status_code, 200)

        self.fechamento.refresh_from_db()
        self.assertTrue(self.fechamento.cancelado)
        self.assertEqual(self.fechamento.cancelado_por, self.gerente)

        # A prova tem que passar pelo endpoint de verdade: e ele que decide
        # qual manager usar (`objects`, sem cancelados, ou `todos`), e so
        # contar direto no ORM nao pegaria a view usando o manager errado.
        listagem = self.client_api.get("/api/v1/fechamentos-caixa/")
        self.assertEqual(listagem.status_code, 200)
        self.assertEqual(lista(listagem), [])

    def test_gerencia_cancela_mesmo_fora_da_janela(self):
        """A janela e da loja. A gerencia confere depois, e e ai que ela ve."""
        FechamentoCaixa.todos.filter(id=self.fechamento.id).update(
            created_at=timezone.now() - timedelta(hours=5)
        )
        self.assertEqual(self.client_api.post(self.url()).status_code, 200)

    def test_cancelar_duas_vezes_nao_reescreve_quem_cancelou(self):
        """O botao pode ser tocado duas vezes num celular ruim."""
        self.client_api.post(self.url())
        self.fechamento.refresh_from_db()
        primeiro_carimbo = self.fechamento.cancelado_em

        self.assertEqual(self.client_api.post(self.url()).status_code, 200)
        self.fechamento.refresh_from_db()
        self.assertEqual(self.fechamento.cancelado_em, primeiro_carimbo)

    def test_funcionario_com_login_nao_cancela(self):
        """Antes disto o DELETE do ModelViewSet deixava o funcionario apagar."""
        funcionario = User.objects.create_user(username="func@x.com", password="123456")
        funcionario.groups.add(Group.objects.get(name="Funcionario"))
        vincular_conta(funcionario, self.conta)

        outro = APIClient()
        outro.force_authenticate(funcionario)
        self.assertEqual(outro.post(self.url()).status_code, 403)

    def test_gerencia_de_outra_conta_nao_cancela(self):
        from app.models import Conta

        alheia = Conta.objects.create(nome="Outro negocio")
        forasteiro = User.objects.create_user(username="alheio@x.com", password="123456")
        forasteiro.groups.add(Group.objects.get(name="Gerente"))
        vincular_conta(forasteiro, alheia)

        outro = APIClient()
        outro.force_authenticate(forasteiro)
        self.assertEqual(outro.post(self.url()).status_code, 404)

    def test_o_turno_aceita_um_lancamento_novo_depois_do_cancelamento(self):
        """E o caso que originou o pedido: mandou errado, cancela, manda certo.

        O relancamento tem que passar pelo POST publico de verdade — e o
        aparelho da loja quem manda, nao o painel da gerencia — para provar
        que a loja consegue mesmo relancar o turno, e nao so que o banco
        aceitaria uma segunda linha.
        """
        self.client_api.post(self.url())

        # Quem relanca e o aparelho da loja, sem login nenhum — nao o painel
        # logado da gerencia que acabou de cancelar. Um client anonimo novo,
        # so com o cookie do aparelho, e o que prova que a LOJA consegue
        # relancar, e nao so que a gerencia (que sempre teria PodeConferirOCaixa
        # de sobra) consegue.
        aparelho = APIClient()
        resposta_acesso = aparelho.post(
            "/api/v1/formulario/acesso/",
            {"codigo": self.conta.codigo_acesso, "apelido": "Celular do balcao"},
            format="json",
        )
        self.assertEqual(resposta_acesso.status_code, 200)

        quem_lanca = Encarregado.objects.create(nome="Marina", conta=self.conta)
        resposta = aparelho.post(
            "/api/v1/fechamentos-caixa/",
            {
                "loja": str(self.fechamento.loja.public_id),
                "lancado_por": str(quem_lanca.public_id),
                "data": DIA_COMUM,
                "periodo": "MANHA",
                "pix": "0.00",
                "cartao": "0.00",
                "dinheiro": "0.00",
                "link_pagamento": "0.00",
                "houve_retirada": False,
                "houve_despesa": False,
                "houve_desperdicio": False,
            },
            format="json",
        )
        self.assertEqual(resposta.status_code, 201, resposta.data)

        self.assertEqual(FechamentoCaixa.objects.count(), 1)
        self.assertEqual(FechamentoCaixa.objects.first().nome_funcionario, "Marina")


class CancelarPelaLojaTests(TestCase):
    """A loja cancela sem login, pelo aparelho, dentro da janela.

    O aparelho entra com o codigo da empresa e recebe um cookie — e o mesmo
    caminho que o formulario ja usa para lancar. O cache e limpo porque o
    throttle do acesso e 5/min e vaza de um teste para o outro.
    """

    def setUp(self):
        from django.core.cache import cache

        cache.clear()
        self.conta = conta_padrao()
        self.fechamento = criar_fechamento()
        self.client_api = APIClient()

    def tearDown(self):
        from django.core.cache import cache

        cache.clear()

    def conectar_aparelho(self, conta=None):
        """Deixa o APIClient com o cookie do aparelho daquela empresa."""
        conta = conta or self.conta
        resposta = self.client_api.post(
            "/api/v1/formulario/acesso/",
            {"codigo": conta.codigo_acesso, "apelido": "Celular do balcao"},
            format="json",
        )
        self.assertEqual(resposta.status_code, 200)

    def url(self, fechamento=None):
        alvo = fechamento or self.fechamento
        return f"/api/v1/fechamentos-caixa/{alvo.public_id}/cancelar/"

    def test_sem_aparelho_conectado_nao_cancela(self):
        # 401 e nao 403: sem o cookie do aparelho a requisicao chega anonima,
        # e o DRF responde "sem credenciais" (401), nao "credenciais
        # insuficientes" (403) — a diferenca so aparece quando se checa o
        # numero certo em vez de aceitar os dois.
        self.assertEqual(self.client_api.post(self.url()).status_code, 401)

    def test_a_loja_cancela_dentro_da_janela_e_sem_carimbar_usuario(self):
        """cancelado_por nulo e o que diz "foi a loja": ela nao tem login."""
        self.conectar_aparelho()

        self.assertEqual(self.client_api.post(self.url()).status_code, 200)

        self.fechamento.refresh_from_db()
        self.assertTrue(self.fechamento.cancelado)
        self.assertIsNone(self.fechamento.cancelado_por)
        self.assertEqual(FechamentoCaixa.objects.count(), 0)

    def test_depois_dos_vinte_minutos_a_loja_nao_cancela_mais(self):
        self.conectar_aparelho()
        FechamentoCaixa.todos.filter(id=self.fechamento.id).update(
            created_at=timezone.now() - timedelta(minutes=21)
        )

        resposta = self.client_api.post(self.url())

        self.assertEqual(resposta.status_code, 403)
        self.fechamento.refresh_from_db()
        self.assertFalse(self.fechamento.cancelado)

    def test_o_aparelho_de_uma_empresa_nao_cancela_da_outra(self):
        """O public_id e imprevisivel, mas nao diz de quem e o lancamento: sem
        o filtro por conta, o id da vizinha bastaria."""
        from app.models import Conta

        alheia = Conta.objects.create(nome="Outro negocio")
        da_alheia = criar_fechamento(
            loja=criar_loja(
                conta=alheia, nome_loja="Loja da outra", cidade="X", endereco="Y"
            )
        )
        self.conectar_aparelho()

        self.assertEqual(self.client_api.post(self.url(da_alheia)).status_code, 404)


class DeleteNaoResponde(TestCase):
    """O DELETE do ModelViewSet estava aberto para funcionario com login, e
    apagava em cascata as despesas e o consumo do turno — o consumo e o que a
    dona desconta no fim do mes. Cancelar substitui isso; apagar de vez segue
    existindo so no /admin."""

    def setUp(self):
        for nome in ("Admin", "Gerente"):
            Group.objects.get_or_create(name=nome)
        self.conta = conta_padrao()
        self.fechamento = criar_fechamento()
        self.gerente = User.objects.create_user(username="ger@x.com", password="123456")
        self.gerente.groups.add(Group.objects.get(name="Gerente"))
        vincular_conta(self.gerente, self.conta)
        self.client_api = APIClient()
        self.client_api.force_authenticate(self.gerente)

    def test_delete_devolve_405(self):
        resposta = self.client_api.delete(
            f"/api/v1/fechamentos-caixa/{self.fechamento.public_id}/"
        )
        self.assertEqual(resposta.status_code, 405)
        self.assertEqual(FechamentoCaixa.todos.count(), 1)
