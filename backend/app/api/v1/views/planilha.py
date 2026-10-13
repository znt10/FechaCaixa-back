"""A exportacao do periodo para planilha.

A dona controlava tudo numa planilha a mao, e continua fechando o mes nela: o
painel responde na tela, mas o desconto em folha e a conferencia com o contador
acontecem no Excel. Este endpoint devolve o periodo inteiro num arquivo so.

Sao cinco abas, e nao uma tabela unica, porque as coisas tem grao diferente:
uma linha de entrada e um turno de uma loja, uma linha de saida e uma despesa
ou uma retirada, uma linha de consumo e uma pessoa num turno, e uma linha de
desperdicio e um item jogado fora. Juntas numa aba so, metade das colunas
ficaria vazia — e alguem arrastaria a soma pela coluna de valor, misturando o
caixa da loja com o lanche de quem trabalha nela, ou com o item que nem foi
vendido. E o mesmo erro que o consumo e o desperdicio separados do caixa
existem para evitar.

Quem baixa escolhe quais abas vem (`?abas=`), e sem escolha vem todas: no dia
25 o consumo vai sozinho para o financeiro, que desconta da folha e nao tem nada
a ver com o caixa da loja.

As abas se chamam pelo tipo do lancamento — Entradas, Saidas, Consumo,
Desperdicio — e nao pela tela de onde vieram: e assim que a empresa fala do
dinheiro.

"Saidas" e o nome do evento, nao do efeito na conta: retirada e despesa saem da
gaveta e VOLTAM para o total (ver FechamentoCaixa.total). So a devolucao nao
volta. O desperdicio nao entra em Saidas: ninguem pagou nada, entao nenhum
dinheiro saiu da gaveta — somar a aba Saidas continua batendo com as colunas
de retirada, despesa e devolucao da aba Entradas.
"""

import unicodedata
from datetime import date

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from rest_framework.exceptions import ValidationError
from rest_framework.permissions import IsAuthenticated
from rest_framework.views import APIView

from app.models import FechamentoCaixa
from app.permissions import PodeConferirOCaixa, get_conta_do_usuario

from .fechamentos import escopo_de_lojas

# Formatos do Excel, e nao texto ja formatado: numero que chega como texto nao
# soma, e data como texto nao ordena — que sao as duas coisas que ela vai fazer
# assim que o arquivo abrir.
DINHEIRO = "#,##0.00"
DIA = "DD/MM/YYYY"
INTEIRO = "0"

CABECALHO = Font(bold=True, color="FFFFFF")
FECHO = Font(bold=True)
FUNDO_DO_CABECALHO = PatternFill("solid", fgColor="1F6F4A")

TURNOS = dict(FechamentoCaixa.Periodo.choices)

# As abas na ordem do arquivo. O nome e o que a tela manda em ?abas= e tambem o
# do metodo que escreve a aba (`_entradas`, `_saidas`...).
ABAS = ("entradas", "saidas", "consumo_por_pessoa", "consumo", "desperdicio")


def _alfabetica(nome):
    """A chave de ordem de um nome de pessoa: sem acento e sem maiuscula.

    Pela ordem crua dos caracteres o "Â" vem depois do "z", e a Ângela iria
    parar no fim da lista — que e uma lista de procura por nome.
    """
    decomposto = unicodedata.normalize("NFKD", nome)
    return "".join(c for c in decomposto if not unicodedata.combining(c)).casefold()


def _turno(fechamento):
    """O turno por extenso.

    O turno do dia inteiro tem dois sentidos. Na empresa de manha e tarde ele
    so acontece no feriado, e escrito "Dia" ninguem sabia de que turno se
    tratava — sai "Feriado", como o domingo sai "Domingo". Na empresa que
    fecha uma vez por dia ele e o fechamento de todo dia, e continua "Dia".
    """
    if (
        fechamento.periodo == FechamentoCaixa.Periodo.DIA
        and fechamento.loja.conta.fechamentos_por_dia != 1
    ):
        return "Feriado"
    return TURNOS.get(fechamento.periodo, fechamento.periodo)


def _aba(planilha, titulo, colunas):
    """Cria a aba com o cabecalho congelado e as larguras ja dadas.

    Congelar a primeira linha e o que faz uma lista de 300 linhas continuar
    legivel na terceira rolagem.
    """
    folha = planilha.create_sheet(titulo)
    folha.append([nome for nome, _ in colunas])
    for coluna, (_, largura) in enumerate(colunas, start=1):
        celula = folha.cell(row=1, column=coluna)
        celula.font = CABECALHO
        celula.fill = FUNDO_DO_CABECALHO
        celula.alignment = Alignment(vertical="center")
        folha.column_dimensions[get_column_letter(coluna)].width = largura
    folha.freeze_panes = "A2"
    return folha


def _formatar(folha, formatos):
    """Aplica o formato de cada coluna nas linhas de dados."""
    for coluna, formato in formatos.items():
        for linha in range(2, folha.max_row + 1):
            folha.cell(row=linha, column=coluna).number_format = formato


def _numero(valor):
    return float(valor or 0)


def _em_blocos(folha, linhas, rotulo_em, somar, nome_do_bloco=str, ordem=None):
    """Escreve as linhas agrupadas, cada bloco fechado por um subtotal.

    `linhas` e uma lista de (bloco, valores) — o bloco e a loja nas abas do
    caixa e a pessoa na do consumo. Ela confere bloco por bloco; sem o
    subtotal, o total de cada um sai de arrastar o mouse pela coluna, e e ai
    que se soma o bloco do vizinho junto.

    O bloco e qualquer chave: a pessoa entra pelo cadastro, e nao pelo nome,
    para dois Bruno continuarem sendo dois descontos. `nome_do_bloco` diz o
    que escrever no subtotal e `ordem` em que ordem os blocos saem.

    Sem nenhuma linha nao escreve nada: um periodo vazio nao pode aparecer com
    uma linha de TOTAL zero, que se le como "a loja nao vendeu" em vez de "nao
    ha nada aqui".

    O rotulo vai na ultima coluna de texto antes dos numeros, que e onde o
    Excel poe subtotal desde sempre — e onde o olho procura.
    """
    if not linhas:
        return

    def fechar(rotulo, bloco):
        fecho = [None] * folha.max_column
        fecho[rotulo_em - 1] = rotulo
        for coluna in somar:
            fecho[coluna - 1] = sum(_numero(valores[coluna - 1]) for valores in bloco)
        folha.append(fecho)
        for celula in folha[folha.max_row]:
            celula.font = FECHO

    for chave in sorted({chave for chave, _ in linhas}, key=ordem):
        bloco = [valores for de_quem, valores in linhas if de_quem == chave]
        for valores in bloco:
            folha.append(valores)
        fechar(f"Total {nome_do_bloco(chave)}", bloco)

    fechar("TOTAL", [valores for _, valores in linhas])


class PlanilhaDoPeriodoView(APIView):
    """GET /api/v1/planilha/?de=&ate=&abas= — o periodo inteiro em .xlsx.

    Mesma permissao da leitura do painel: quem pode conferir o caixa pode
    levar o periodo para a planilha. O escopo sai de `escopo_de_lojas`, o
    mesmo do painel, para que a exportacao nao seja uma segunda regra de
    isolamento entre empresas — divergir dela e como um vazamento comeca.
    """

    permission_classes = [IsAuthenticated, PodeConferirOCaixa]

    def get(self, request):
        hoje = timezone.localdate()
        abas = self._abas(request.query_params.get("abas"))
        de = self._data(request.query_params.get("de")) or hoje
        ate = self._data(request.query_params.get("ate")) or hoje

        lojas = escopo_de_lojas(request)
        fechamentos = list(
            FechamentoCaixa.objects.filter(loja__in=lojas, data__gte=de, data__lte=ate)
            # A conta vem junto porque o nome do turno do dia inteiro depende
            # dela (ver _turno): sem isto, uma consulta por linha.
            .select_related("loja__conta")
            .prefetch_related(
                "retiradas__responsavel", "consumos__encarregado", "despesas",
                "desperdicios__salgado__categoria",
            )
            .order_by("data", "loja__nome_loja", "periodo")
        )

        planilha = Workbook()
        # O Workbook ja nasce com uma aba, e ela nao e nenhuma das nossas:
        # sem remover, o arquivo abre numa folha vazia.
        planilha.remove(planilha.active)

        for aba in abas:
            getattr(self, f"_{aba}")(planilha, fechamentos)

        resposta = HttpResponse(
            content_type=(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
        )
        resposta["Content-Disposition"] = (
            f'attachment; filename="{self._nome_do_arquivo(request, de, ate)}"'
        )
        # A planilha e do periodo pedido, e o periodo de hoje muda a cada
        # lancamento: guardada, ela mostraria um mes que ja mudou.
        resposta["Cache-Control"] = "no-store"
        planilha.save(resposta)
        return resposta

    @staticmethod
    def _abas(texto):
        """As abas pedidas, na ordem do arquivo — e nao na de quem marcou.

        Sem o parametro vem todas, que e o arquivo de sempre. Nome que nao e
        de aba nenhuma e ignorado; se nao sobrar nenhuma, e 400, porque um
        .xlsx sem aba nao abre no Excel.
        """
        if not texto:
            return ABAS
        pedidas = {nome.strip() for nome in texto.split(",")}
        abas = [aba for aba in ABAS if aba in pedidas]
        if not abas:
            raise ValidationError({"abas": "Escolha ao menos uma aba da planilha."})
        return abas

    @staticmethod
    def _data(texto):
        try:
            return date.fromisoformat(texto) if texto else None
        except ValueError:
            # Data ilegivel vira "hoje" em vez de 500: o parametro vem da URL,
            # e a tela nunca manda outra coisa.
            return None

    def _nome_do_arquivo(self, request, de, ate):
        conta = get_conta_do_usuario(request.user)
        empresa = conta.slug if conta else "todas-as-empresas"
        periodo = de.isoformat() if de == ate else f"{de.isoformat()}-a-{ate.isoformat()}"
        return f"fechacaixa-{empresa}-{periodo}.xlsx"

    # ------------------------------------------------------------------ abas

    def _entradas(self, planilha, fechamentos):
        """Um turno de uma loja por linha: a conta fechando na horizontal.

        A ordem das colunas e a da conta, e nao a da tela: as seis parcelas
        primeiro (PIX, cartao, dinheiro, link, retirada, despesa) e o Total
        logo depois delas. Assim quem arrasta a soma pela linha chega no mesmo
        numero que a coluna Total ja traz — que e como a dona confere.

        Retirada e despesa SOMAM: `dinheiro` e o que sobrou na gaveta, e as
        duas sairam dali depois da venda. A devolucao fica na ponta, depois do
        Total, justamente por nao entrar nele — ela cancelou a venda junto, e
        ja se descontou sozinha no dinheiro contado.

        Nao ha mais "Total liquido": com nada subtraindo, ele repetia o Total.
        """
        folha = _aba(
            planilha,
            "Entradas",
            [
                ("Data", 12), ("Loja", 22), ("Turno", 10), ("Quem lançou", 24),
                ("PIX", 12), ("Cartão", 12), ("Dinheiro", 12), ("Link", 12),
                ("Retirada", 12), ("Despesa", 12), ("Total", 13),
                ("Devolução", 12), ("Conferido", 11),
            ],
        )
        linhas = [
            (f.loja.nome_loja, [
                f.data,
                f.loja.nome_loja,
                _turno(f),
                f.nome_funcionario,
                _numero(f.pix), _numero(f.cartao), _numero(f.dinheiro),
                _numero(f.link_pagamento),
                _numero(f.valor_retirado) if f.houve_retirada else 0,
                _numero(f.total_das_despesas),
                _numero(f.total),
                _numero(f.devolucao_valor) if f.houve_devolucao else 0,
                "sim" if f.conferido else "não",
            ])
            for f in fechamentos
        ]
        # Rotulo em "Quem lancou", a ultima coluna de texto antes do dinheiro.
        _em_blocos(folha, linhas, rotulo_em=4, somar=range(5, 13))
        _formatar(folha, {1: DIA, **{c: DINHEIRO for c in range(5, 13)}})

    def _saidas(self, planilha, fechamentos):
        """O dinheiro que saiu da gaveta: despesa, retirada e devolucao.

        Uma linha por evento — o turno que comprou gas e agua e mandou dinheiro
        para o dono vira tres linhas, com donos diferentes, que a gerencia
        cobra de gente diferente.

        Consumo NAO entra aqui, e e o ponto da separacao: ninguem pagou na
        hora, entao ele nao saiu do caixa. Somar esta aba tem que dar o mesmo
        que as colunas de retirada, despesa e devolucao da aba Entradas.

        "Saiu da gaveta" nao quer dizer "saiu do total": retirada e despesa
        voltam para o total na aba Entradas, e so a devolucao nao volta. Esta
        aba responde para onde o dinheiro foi, nao quanto a loja vendeu.
        """
        folha = _aba(
            planilha,
            "Saídas",
            [
                ("Data", 12), ("Loja", 22), ("Turno", 10), ("Tipo", 11),
                ("Descrição", 34), ("Valor", 13), ("Quem lançou", 24),
            ],
        )
        linhas = []
        for f in fechamentos:
            comum = [f.data, f.loja.nome_loja, _turno(f)]
            for despesa in sorted(f.despesas.all(), key=lambda d: d.descricao):
                if _numero(despesa.valor) <= 0:
                    continue
                linhas.append((f.loja.nome_loja, comum + [
                    "Despesa", despesa.descricao,
                    _numero(despesa.valor), f.nome_funcionario,
                ]))
            # Uma linha por pessoa: o dono e a socia que retiraram no mesmo
            # turno sao cobrados separados no fim do mes. Sem linhas, cai no
            # resumo do fechamento — e o que sobra de um lancamento gravado
            # por fora do serializer, e a soma continua tendo que bater com a
            # coluna Retirada da aba Entradas.
            retiradas = [
                (r.responsavel.nome if r.responsavel else "Retirada", r.valor)
                for r in f.retiradas.all()
            ]
            if not retiradas and f.houve_retirada:
                retiradas = [(
                    f.responsavel_retirada.nome if f.responsavel_retirada else "Retirada",
                    f.valor_retirado,
                )]
            for quem, valor in retiradas:
                if _numero(valor) <= 0:
                    continue
                linhas.append((f.loja.nome_loja, comum + [
                    "Retirada", quem, _numero(valor), f.nome_funcionario,
                ]))
            if f.houve_devolucao and _numero(f.devolucao_valor) > 0:
                linhas.append((f.loja.nome_loja, comum + [
                    "Devolução", "Dinheiro devolvido ao cliente",
                    _numero(f.devolucao_valor), f.nome_funcionario,
                ]))
        # Rotulo em "Descricao", a ultima coluna de texto antes do valor.
        _em_blocos(folha, linhas, rotulo_em=5, somar=[6])
        _formatar(folha, {1: DIA, 6: DINHEIRO})

    def _consumo_por_pessoa(self, planilha, fechamentos):
        """O numero do desconto: uma linha por pessoa, com o total do periodo.

        A pessoa nao e fixa numa loja — hoje esta na Lapa, amanha no Limao —
        entao a coluna de lojas existe para mostrar que ela circulou, e nao
        para agrupar.
        """
        folha = _aba(
            planilha,
            "Consumo por pessoa",
            [("Pessoa", 32), ("Total", 13), ("Lançamentos", 13), ("Lojas", 40)],
        )

        por_pessoa = {}
        for f in fechamentos:
            for consumo in f.consumos.all():
                if _numero(consumo.valor) <= 0:
                    continue
                linha = por_pessoa.setdefault(
                    consumo.encarregado_id,
                    {"nome": consumo.encarregado.nome, "total": 0.0,
                     "vezes": 0, "lojas": []},
                )
                linha["total"] += _numero(consumo.valor)
                linha["vezes"] += 1
                if f.loja.nome_loja not in linha["lojas"]:
                    linha["lojas"].append(f.loja.nome_loja)

        # A a Z, e nao do maior para o menor: esta aba e uma lista de procura.
        # Ela abre aqui para achar UMA pessoa e ver quanto descontar da folha —
        # por nome se acha na hora, por valor se percorre a lista inteira. Quem
        # consumiu mais e a pergunta da tela, e la ela continua respondida.
        # O id desempata os xaras na mesma ordem da aba Consumo.
        for _, linha in sorted(
            por_pessoa.items(), key=lambda par: (_alfabetica(par[1]["nome"]), par[0])
        ):
            folha.append([
                linha["nome"], linha["total"], linha["vezes"],
                ", ".join(sorted(linha["lojas"])),
            ])
        _formatar(folha, {2: DINHEIRO})

    def _consumo(self, planilha, fechamentos):
        """O detalhe por tras do total: e o que ela mostra quando alguem
        contesta o desconto.

        Agrupado por pessoa, de A a Z, com os lancamentos de cada uma em ordem
        de data e um subtotal no fim. Em ordem de data, o consumo de uma pessoa
        ficava espalhado pelo mes inteiro; assim o subtotal de cada uma e o
        mesmo numero da aba Consumo por pessoa, e as duas vao juntas para o
        financeiro.
        """
        folha = _aba(
            planilha,
            "Consumo",
            [
                ("Data", 12), ("Loja", 22), ("Turno", 10), ("Pessoa", 32),
                ("Valor", 13), ("Quem lançou", 24),
            ],
        )
        # Os fechamentos ja vem por data, loja e turno: e essa a ordem que
        # sobra dentro de cada pessoa.
        linhas = [
            ((consumo.encarregado.nome, consumo.encarregado_id), [
                f.data, f.loja.nome_loja, _turno(f),
                consumo.encarregado.nome, _numero(consumo.valor),
                f.nome_funcionario,
            ])
            for f in fechamentos
            for consumo in f.consumos.all()
        ]
        # Rotulo em "Pessoa", a ultima coluna de texto antes do valor.
        _em_blocos(
            folha, linhas, rotulo_em=4, somar=[5],
            nome_do_bloco=lambda pessoa: pessoa[0],
            ordem=lambda pessoa: (_alfabetica(pessoa[0]), pessoa[1]),
        )
        _formatar(folha, {1: DIA, 5: DINHEIRO})

    def _desperdicio(self, planilha, fechamentos):
        """O que foi para o lixo, uma linha por item.

        Aba propria e nao uma linha em Saidas: aquela aba e dinheiro que saiu
        da gaveta, e declara que a soma dela bate com as colunas de retirada,
        despesa e devolucao da aba Entradas. O desperdicio nao tem valor —
        ninguem pagou nada — e uma quantidade na coluna Valor seria somada com
        reais na primeira vez que alguem arrastasse o mouse.

        Por isso a coluna Quantidade e INTEIRO e nunca DINHEIRO: o formato e o
        que avisa o olho de que aquilo nao e R$.
        """
        folha = _aba(
            planilha,
            "Desperdício",
            [
                ("Data", 12), ("Loja", 22), ("Turno", 10), ("Categoria", 18),
                ("Item", 26), ("Quantidade", 13), ("Quem lançou", 24),
            ],
        )
        linhas = []
        for f in fechamentos:
            for perda in sorted(
                f.desperdicios.all(), key=lambda d: d.salgado.nome
            ):
                if perda.quantidade <= 0:
                    continue
                linhas.append((f.loja.nome_loja, [
                    f.data,
                    f.loja.nome_loja,
                    _turno(f),
                    perda.salgado.categoria.nome,
                    perda.salgado.nome,
                    perda.quantidade,
                    f.nome_funcionario,
                ]))
        # Rotulo em "Item", a ultima coluna de texto antes do numero.
        _em_blocos(folha, linhas, rotulo_em=5, somar=[6])
        _formatar(folha, {1: DIA, 6: INTEIRO})
